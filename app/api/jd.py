# app/api/jd.py

from fastapi import APIRouter
from app.agents.jd_generator import generate_jd
from app.agents.jd_clarifier import generate_clarifying_questions
from app.agents.jd_chatbot import refine_jd
from app.agents.profile_builder import build_profile
import json
import os
import gspread
import pandas as pd
from google.oauth2.service_account import Credentials

router = APIRouter(
    prefix="/jd",
    tags=["Job Description"]
)


@router.get("/roles")
def get_roles():
    """Fetch available job roles from Google Sheets."""
    from pathlib import Path
    import re

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets.readonly",
        "https://www.googleapis.com/auth/drive.readonly",
    ]

    # 1) Try environment variable
    service_account_json = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "")

    # 2) Try reading from .streamlit secrets files
    if not service_account_json:
        project_root = Path(__file__).resolve().parents[2]
        for secrets_name in [".secrets", "secrets.toml"]:
            secrets_path = project_root / ".streamlit" / secrets_name
            if secrets_path.exists():
                raw = secrets_path.read_text(encoding="utf-8")
                # Extract the JSON between triple quotes
                match = re.search(r"GOOGLE_SERVICE_ACCOUNT_JSON\s*=\s*'''(.*?)'''", raw, re.DOTALL)
                if not match:
                    match = re.search(r'GOOGLE_SERVICE_ACCOUNT_JSON\s*=\s*"""(.*?)"""', raw, re.DOTALL)
                if match:
                    service_account_json = match.group(1).strip()
                    break

    if not service_account_json:
        return []

    service_account_info = json.loads(service_account_json)
    creds = Credentials.from_service_account_info(service_account_info, scopes=scopes)
    client = gspread.authorize(creds)

    SPREADSHEET_ID = "1SpNGsY707CaY6i06knI9F2HJdtAcHxGKq8IjAb17oWo"
    sheet = client.open_by_key(SPREADSHEET_ID).sheet1
    df = pd.DataFrame(sheet.get_all_records())
    df.columns = [c.strip().lower() for c in df.columns]

    result = []
    for _, row in df.iterrows():
        result.append({
            "role": row.get("job title ( example: ai engineer, sales executive, hr manager)", ""),
            "department": row.get("in which department (ex. marketing, tech etc.)", ""),
            "location": row.get("location", ""),
            "employment_type": row.get("employment type ( full-time / contract / internship )", "Full-time"),
            "travel_required": row.get("does this role require travel?", ""),
            "work_mode": row.get("work mode", ""),
            "key_responsibilities": row.get("key responsibilities  ( list 4–6 things this person will actually do)", ""),
            "reporting_to": row.get("reporting to (example: tech lead, sales manager)", ""),
            "new_or_scaling": row.get("is this role building something new or scaling an existing function?", ""),
            "must_have_skills": row.get("top 3 skills this role must have", ""),
            "other_skills": row.get("other skills ( example: python, excel, communication )", ""),
            "minimum_education": row.get("minimum education required", ""),
            "experience": row.get("minimum experience required", ""),
            "urgency": row.get("how urgent is this hire?", ""),
            "salary": row.get("salary range (optional)", ""),
        })

    return result

@router.post("/clarify")
def clarify_jd_api(payload: dict):
    """Agent 1: Generate clarifying questions from form data (no draft needed)."""
    questions = generate_clarifying_questions(form_data=payload)

    if isinstance(questions, str):
        questions = json.loads(questions)

    return {"questions": questions}


@router.post("/profile")
def profile_builder_api(payload: dict):
    """Agent 2: Build ideal candidate profile."""
    profile = build_profile(
        form_data=payload.get("form_data", {}),
        clarification_answers=payload.get("answers", [])
    )
    return {"profile": profile}


@router.post("/generate")
def generate_jd_api(payload: dict):
    """Agent 3: Generate JD from profile + form data."""
    jd = generate_jd(
        form_data=payload.get("form_data", payload),
        profile=payload.get("profile", None)
    )
    return {"jd": jd}


@router.post("/refine")
def refine_jd_api(payload: dict):
    """Agent 4: Refine JD based on user instruction."""
    updated_jd = refine_jd(
        current_jd=payload["jd"],
        instruction=payload["instruction"],
        role=payload.get("role", ""),
        session_id=payload.get("session_id", "")
    )
    return {"jd": updated_jd}


@router.post("/export-docx")
def export_jd_docx(payload: dict):
    """Export JD as a formatted .docx file."""
    import io
    import re
    from fastapi.responses import StreamingResponse
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    jd_text = payload.get("jd", "")
    role = payload.get("role", "Job Description")

    doc = Document()

    # ── Page margins ──
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Default font ──
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    # ── Parse markdown lines ──
    lines = jd_text.split("\n")

    for line in lines:
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            continue

        # Heading 1: # Title
        if stripped.startswith("# ") and not stripped.startswith("## "):
            heading = stripped[2:].strip()
            p = doc.add_heading(heading, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
            continue

        # Heading 2: ## Section
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            p = doc.add_heading(heading, level=2)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x2d, 0x2d, 0x5e)
                run.font.size = Pt(13)
            continue

        # Heading 3: ### Sub-section
        if stripped.startswith("### "):
            heading = stripped[4:].strip()
            p = doc.add_heading(heading, level=3)
            for run in p.runs:
                run.font.size = Pt(12)
            continue

        # Bullet points: - item or * item
        if re.match(r"^[-*•]\s+", stripped):
            text = re.sub(r"^[-*•]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            # Handle bold markers **text**
            _add_formatted_run(p, text)
            continue

        # Numbered list: 1. item
        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _add_formatted_run(p, text)
            continue

        # Horizontal rules
        if re.match(r"^[-─═]{3,}$", stripped):
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _add_formatted_run(p, stripped)

    # ── Save to buffer ──
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = f"{role.replace(' ', '_')}_JD.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def _add_formatted_run(paragraph, text: str):
    """Add text to a paragraph, converting **bold** markers to bold runs."""
    import re
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

